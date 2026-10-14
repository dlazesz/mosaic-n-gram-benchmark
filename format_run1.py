import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule


def parse_measurements(input_file):
    """Processing of measurement data, using only the "real" value"""

    pattern = re.compile(r'Starting measurements:\s+'
                         r'(\w+)\s+\(([^)]+)\)\s+'
                         r'(\d+)-gram,\s+'
                         r'(fixed-length|variable-length)\s+fields,\s+'
                         r'padding(?:\s+words)?\s+([\d ]+)\s+lines',
                         re.IGNORECASE)

    measurements = []
    current = None

    with open(input_file, encoding='UTF-8') as fh:
        for line in fh:
            line = line.strip()

            match = pattern.search(line)
            if match is not None:
                language, method, ngram, fields, lines = match.groups()

                current = {'Key': f'{language.upper()} ({method})',
                           'Language': language.upper(),
                           'Method': method,
                           'N-gram size': int(ngram),
                           'Fields': ('fixed' if fields.lower() == 'fixed-length' else 'variable'),
                           'No. of lines': int(lines.replace(' ', '')),
                           'Real (seconds)': None}

                measurements.append(current)

            elif current is not None:
                time_match = re.match(r'real\s*:\s*([\d.,]+)', line, re.IGNORECASE)
                if time_match is not None:
                    current['Real (seconds)'] = float(time_match.group(1).replace(',', '.'))

    if len(measurements) == 0:
        raise ValueError('No measurement found in the input file')

    missing = [index for index, item in enumerate(measurements, start=1) if item['Real (seconds)'] is None]
    if len(missing) > 0:
        raise ValueError(f'Missing real value in the following measurements: {missing}')

    return measurements


def format_sheet(ws):
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions

    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill(fill_type='solid', fgColor='1F4E78')
        cell.alignment = Alignment(horizontal='center')

    ws.row_dimensions[1].height = 25

    for column in ws.columns:
        max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column)

        letter = get_column_letter(column[0].column)
        ws.column_dimensions[letter].width = min(max(max_length + 2, 17), 48)

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, (int, float)):
                cell.number_format = '0.00000'


def create_comparison_sheet(wb, measurements, fields, ngram, line_count):
    """One sheet for a single combination of parameters"""

    sheet_name = f'{fields}_{ngram}gram_{line_count}'
    ws = wb.create_sheet(title=sheet_name[:31])

    filtered = [item for item in measurements
                if item['Fields'] == fields and item['N-gram size'] == ngram and item['No. of lines'] == line_count]

    # Language and method together form the unique key
    values = {item['Key']: item['Real (seconds)'] for item in filtered}

    # An ordered list of implementations
    keys = sorted(values.keys())

    ws.append(['Language and Method', 'Real (seconds)'])

    for key in keys:
        ws.append([key, values[key]])

    # Shorter running times are shown in green, longer ones in red
    if ws.max_row >= 2:
        ws.conditional_formatting.add(f'B2:B{ws.max_row}',
                                      ColorScaleRule(start_type='min',
                                                     start_color='63BE7B',
                                                     mid_type='percentile',
                                                     mid_value=50,
                                                     mid_color='FFEB84',
                                                     end_type='max',
                                                     end_color='F8696B'))

    format_sheet(ws)


def create_excel(input_file, output_file):
    measurements = parse_measurements(input_file)

    wb = Workbook()
    active = wb.active
    if active is not None:
        wb.remove(active)

    # Each different configuration is assigned its own worksheet
    configurations = sorted({(item['Fields'], item['N-gram size'], item['No. of lines']) for item in measurements})
    for fields, ngram, line_count in configurations:
        create_comparison_sheet(wb, measurements, fields, ngram, line_count)

    wb.save(output_file)

    print(f'The file has been created: {output_file}')
    print(f'Processed measurements: {len(measurements)}')
    print(f'Number of worksheets: {len(configurations)}')


if __name__ == '__main__':
    if len(sys.argv) != 3:
        print(f'Usage: python {Path(sys.argv[0]).name} measurements.txt results.xlsx')
        sys.exit(1)

    create_excel(sys.argv[1], sys.argv[2])
